from datetime import datetime, timezone
from io import BytesIO

from openpyxl import load_workbook

from .base import FileParser
from .metadata import ParsedMetadata


class XLSXParser(FileParser):
    """Parse XLSX files and extract spreadsheet metadata."""

    def parse(
        self,
        file_name: str,
        mime_type: str,
        file_content: bytes,
    ) -> ParsedMetadata:
        """Parse an XLSX file and return its metadata."""

        workbook = load_workbook(
            BytesIO(file_content),
            read_only=True,
            data_only=True,
        )

        try:
            sheet_count = len(workbook.sheetnames)

            properties = workbook.properties

            title = properties.title or None
            author = properties.creator or None

            return ParsedMetadata(
                file_name=file_name,
                file_type="XLSX",
                mime_type=mime_type,
                file_size=len(file_content),
                sheet_count=sheet_count,
                title=title,
                author=author,
                created_at=datetime.now(timezone.utc),
                parsed_status="SUCCESS",
            )
        finally:
            workbook.close()